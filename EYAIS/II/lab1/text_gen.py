import os
from docx import Document
from sklearn.datasets import fetch_20newsgroups
from tqdm import tqdm

def save_all_newsgroups_to_docx(out_dir="news", per_category=50):
    os.makedirs(out_dir, exist_ok=True)

    print("Загружаю датасет 20 Newsgroups (это может занять время при первом запуске)...")
    newsgroups = fetch_20newsgroups(
        subset="train",
        remove=("headers", "footers", "quotes")
    )

    categories = newsgroups.target_names
    counts = {cat: 0 for cat in categories}
    total_files = per_category * len(categories)

    print(f"Сохраняем {total_files} файлов ({per_category} на каждую из {len(categories)} категорий)...")

    with tqdm(total=total_files, desc="Создание файлов") as pbar:
        for i, text in enumerate(newsgroups.data):
            cat = categories[newsgroups.target[i]]
            if counts[cat] >= per_category:
                continue  # категорию уже набрали

            counts[cat] += 1

            # Создаём docx
            doc = Document()
            doc.add_heading(f"{cat} — статья {counts[cat]}", level=1)
            doc.add_paragraph(f"Категория: {cat}\n")
            doc.add_paragraph(text)

            # Имя файла
            file_name = f"{cat.replace('.', '_')}_{counts[cat]}.docx"
            file_path = os.path.join(out_dir, file_name)
            doc.save(file_path)

            pbar.update(1)  # обновляем прогресс-бар

            if all(c >= per_category for c in counts.values()):
                break

    print(f"✅ Готово! Файлы сохранены в папку '{out_dir}'")

# Пример использования
save_all_newsgroups_to_docx(out_dir="news", per_category=50)
